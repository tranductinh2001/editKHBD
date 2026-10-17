# Beeknoee AI Client for KHBD Pro (edit_KHBD)
import os
import re
import json
import requests
import docx

DEFAULT_API_KEY = "sk-bee-73e9d744f8ef96f61278f28f4acfc8e44b69ee94c9cf6714720044f8c91c8502"
DEFAULT_BASE_URL = "https://platform.beeknoee.com/api/v1"
DEFAULT_MODEL = "gemini-2.5-flash-lite"

class BeeknoeeClient:
    def __init__(self, api_key: str = None, base_url: str = None, model: str = None):
        self.api_key = api_key or os.environ.get("BEE_API_KEY") or DEFAULT_API_KEY
        self.base_url = (base_url or os.environ.get("BEE_BASE_URL") or DEFAULT_BASE_URL).rstrip("/")
        self.model = model or os.environ.get("AI_MODEL") or DEFAULT_MODEL

    def chat(self, prompt: str, system_prompt: str = None, temperature: float = 0.2, max_tokens: int = 4096) -> str:
        messages = []
        if system_prompt:
            messages.append({"role": "system", "content": system_prompt})
        messages.append({"role": "user", "content": prompt})

        payload = {
            "model": self.model,
            "messages": messages,
            "temperature": temperature,
            "max_tokens": max_tokens
        }

        resp = requests.post(
            f"{self.base_url}/chat/completions",
            headers={"Authorization": f"Bearer {self.api_key}", "Content-Type": "application/json"},
            json=payload,
            timeout=90
        )

        if resp.status_code == 200:
            data = resp.json()
            return data["choices"][0]["message"]["content"]
        else:
            raise Exception(f"Beeknoee API Error: {resp.status_code} - {resp.text}")

    def analyze_khbd(self, doc_text: str, dispatch_type: str = "CV_5512") -> dict:
        system_prompt = (
            "Ban la chuyen gia su pham Bo Giao Duc Viet Nam tham dinh Ke hoach bai day (KHBD) "
            f"theo chuan Cong van {dispatch_type.replace('CV_', '')} va Rubric 100 diem. "
            "Chi tra ve duy nhat mot khoi JSON hop le."
        )

        prompt = f'''Hay tham dinh giao an sau theo Rubric 100 diem:

## NOI DUNG GIAO AN:
{doc_text[:6000]}

## RUBRIC 100 DIEM:
1. MUC TIEU / YEU CAU CAN DAT (15 diem): Kien thuc, Nang luc chung, Nang luc dac thu, Pham chat.
2. CHUOI CAC HOAT DONG (35 diem): Du 4 pha (Khoi dong, Hinh thanh kien thuc, Luyen tap, Van dung), thoi gian phu hop.
3. HOAT DONG GIAO VIEN (15 diem): 4 buoc su pham (Chuyen giao, Thuc hien, Bao cao, Ket luan/Nhan xet), su dung CNTT/AI.
4. HOAT DONG HOC SINH (15 diem): HS chu dong lam viec nhom, thao luan, tao ra san pham hoc tap.
5. THIET BI & HOC LIEU (10 diem): Phu hop, tach ro phan GV va HS chuan bi.
6. KIEM TRA & DANH GIA (10 diem): Cong cu danh gia, rubric, san pham hoc tap ro rang.

YEU CAU TRA VE JSON CO DINH DANG CHINH XAC NHU SAU:
{{
  "total_score": 78,
  "grade": "B",
  "scores": {{
    "muc_tieu": {{"score": 12, "max": 15, "comment": "..."}},
    "hoat_dong": {{"score": 28, "max": 35, "comment": "..."}},
    "gv_action": {{"score": 12, "max": 15, "comment": "..."}},
    "hs_action": {{"score": 11, "max": 15, "comment": "..."}},
    "phuong_tien": {{"score": 8, "max": 10, "comment": "..."}},
    "danh_gia": {{"score": 7, "max": 10, "comment": "..."}}
  }},
  "strengths": ["...", "..."],
  "weaknesses": ["...", "..."],
  "suggestions": ["...", "..."],
  "nls_evaluation": "Góp ý về Năng lực số...",
  "cv5512_compliance": "Đánh giá mức độ tuân thủ 4 bước Công văn 5512..."
}}
'''

        raw = self.chat(prompt, system_prompt=system_prompt, temperature=0.1)
        clean = re.sub(r"^```json\s*", "", raw.strip(), flags=re.MULTILINE)
        clean = re.sub(r"^```\s*", "", clean, flags=re.MULTILINE)
        clean = clean.strip()
        
        try:
            return json.loads(clean)
        except Exception:
            m = re.search(r"\{.*\}", clean, re.DOTALL)
            if m:
                return json.loads(m.group(0))
            raise Exception("Khong the phan tich ket qua JSON tu AI: " + raw[:200])


    def extract_lesson_metadata(self, text: str) -> dict:
        prompt = f"""Trich xuat thong tin bai day tu doan dau giao an sau:
{text[:2500]}

Tra ve JSON:
{
  "bai_num": "1",
  "bai_title": "TEN BAI DAY",
  "chude": "TEN CHU DE",
  "tuan": "1",
  "tiet": "1, 2",
  "detected_grade": "9",
  "detected_subject": "Tin hoc"
}
"""
        try:
            raw = self.chat(prompt, system_prompt="Ban la he thong trich xuat thong tin giao an. Chi tra ve JSON duy nhat.", temperature=0.1)
            clean = re.sub(r"^```json\s*", "", raw.strip(), flags=re.MULTILINE)
            clean = re.sub(r"^```\s*", "", clean, flags=re.MULTILINE).strip()
            m = re.search(r"\{.*\}", clean, re.DOTALL)
            if m:
                import json
                return json.loads(m.group(0))
        except Exception as e:
            print(f"AI extract metadata error: {e}")
        return {}

    def analyze_lesson_title_with_ppct(self, text: str, ppct_items: list = None, filename: str = "") -> dict:
        """
        Dùng Beekone AI phân tích đoạn đầu giáo án và đối chiếu PPCT để sửa và chuẩn hoá tên bài chính xác 100%.
        """
        ppct_context = ""
        if ppct_items:
            # Lấy danh sách tóm tắt tối đa 50 bài/tiết để AI đối chiếu
            simplified = []
            for it in ppct_items[:50]:
                simplified.append({
                    "stt": it.get("stt"),
                    "tuan": it.get("tuan"),
                    "tiet": it.get("tiet"),
                    "ten_bai": it.get("ten_bai"),
                    "ghi_chu": it.get("ghi_chu")
                })
            ppct_context = f"\n## DANH SÁCH BÀI HỌC TRONG PHÂN PHỐI CHƯƠNG TRÌNH (PPCT):\n{json.dumps(simplified, ensure_ascii=False, indent=2)}\n"

        prompt = f"""Bạn là chuyên gia sư phạm thẩm định giáo án và đối chiếu Phân Phối Chương Trình (PPCT).
Nhiệm vụ: Phân tích đoạn đầu của giáo án Word dưới đây và danh sách PPCT để xác định chính xác:
1. Số thứ tự bài (bai_num): Số nguyên hoặc số kèm ký tự, ví dụ "1", "2", "3", "11". Nếu là bài ôn tập/kiểm tra thì ghi rõ "Ôn tập" hoặc "Kiểm tra".
2. Tên bài chuẩn (bai_title): Tên bài học chuẩn theo Sách Giáo Khoa / PPCT (viết hoa chữ cái đầu hoặc đúng danh pháp SGK, TUYỆT ĐỐI KHÔNG chứa tiền tố "Bài 1.", "Bài 1:", "Tiết 1+2:", ...).
3. Tên chủ đề (chude): Ví dụ "CHỦ ĐỀ 1: MÁY TÍNH VÀ CỘNG ĐỒNG".
4. Tuần dạy (tuan): Chuỗi số tuần, ví dụ "1", "2".
5. Tiết dạy (tiet): Chuỗi tiết, ví dụ "1, 2" hoặc "3" hoặc "4, 5".
6. Mục khớp trong PPCT (matched_ppct): Chọn 1 item trong danh sách PPCT nếu có, hoặc null nếu không khớp.
7. Giải thích ngắn gọn lý do nhận diện (explanation).

Tên file hiện tại: {filename}

## ĐOẠN ĐẦU NỘI DUNG GIÁO ÁN:
{text[:3500]}
{ppct_context}

## QUY ĐỊNH BẮT BUỘC:
- Trả về DUY NHẤT một khối JSON hợp lệ, không bọc text giải thích bên ngoài.
- Cấu trúc JSON:
{{
  "bai_num": "1",
  "bai_title": "Thế giới kĩ thuật số",
  "chude": "CHỦ ĐỀ 1: MÁY TÍNH VÀ CỘNG ĐỒNG",
  "tuan": "1",
  "tiet": "1, 2",
  "matched_ppct": {{
    "stt": 1,
    "tuan": 1,
    "tiet": 1,
    "ten_bai": "Bài 1. Thế giới kĩ thuật số"
  }},
  "confidence": "high",
  "explanation": "Đoạn đầu giáo án và mục tiêu nêu rõ về thiết bị kĩ thuật số, tương ứng Bài 1 trong PPCT."
}}
"""
        system_prompt = (
            "Bạn là trợ lý AI chuyên gia phân tích bố cục giáo án và chuẩn hoá tên bài học theo PPCT của Bộ GD&ĐT Việt Nam. "
            "Chỉ trả về JSON hợp lệ."
        )

        try:
            raw = self.chat(prompt, system_prompt=system_prompt, temperature=0.1)
            clean = re.sub(r"^```json\s*", "", raw.strip(), flags=re.MULTILINE)
            clean = re.sub(r"^```\s*", "", clean, flags=re.MULTILINE).strip()
            m = re.search(r"\{.*\}", clean, re.DOTALL)
            if m:
                return json.loads(m.group(0))
        except Exception as e:
            print(f"Beekone AI analyze title error: {e}")
            return {
                "error": str(e),
                "confidence": "failed"
            }
        return {}

    def generate_nls_objectives(self, subject: str, grade: str, lesson_title: str) -> str:
        prompt = f"""Soan 2-3 muc tieu ve Nang Luc So (NLS) chuan theo khung nang luc so Bo GD&DT cho:
Mon hoc: {subject}
Khoi lop: {grade}
Bai day: {lesson_title}

Dinh dang:
- Nang luc so (NLS 1.1, NLS 2.2): [Mo ta cu the hanh vi hoc sinh thuc hien duoc gan voi bai hoc]
Tra ve truc tiep cac gach dau dong bang tieng Viet co dau chuan xac.
"""
        try:
            return self.chat(prompt, system_prompt="Ban la chuyen gia Bo GD&DT ve Nang luc so trong chuong trinh GDPT 2018.", temperature=0.2).strip()
        except Exception as e:
            print(f"AI NLS error: {e}")
            return "- Năng lực số (NLS 1.1, NLS 2.3): Học sinh biết tìm kiếm, xử lý và chia sẻ thông tin an toàn trên môi trường số; sử dụng phần mềm hỗ trợ học tập và tạo sản phẩm số cá nhân."

def extract_docx_text(file_path: str) -> str:
    doc = docx.Document(file_path)
    paragraphs_text = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    table_text = []
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                table_text.append(" | ".join(row_cells))
    return "\n".join(paragraphs_text + table_text)
