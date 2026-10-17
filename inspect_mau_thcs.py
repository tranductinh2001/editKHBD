import sys
import io
import docx

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

doc = docx.Document('Mau KHBD THCS.docx')
print('=== PARAGRAPHS IN Mau KHBD THCS.docx ===')
print('Total paragraphs:', len(doc.paragraphs))
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        runs_info = []
        for r in p.runs:
            f = r.font
            sz = f.size.pt if f.size else None
            runs_info.append(f"{r.text}(sz={sz},b={f.bold},i={f.italic})")
        print(f"P{i:03d} [style={p.style.name}, align={p.alignment}]: {p.text}")
        print(f"     runs: {' | '.join(runs_info)}")

print(f'\nTotal tables: {len(doc.tables)}')
for t_idx, t in enumerate(doc.tables):
    print(f'\n--- TABLE {t_idx} ({len(t.rows)}x{len(t.columns)}) ---')
    for r_idx, r in enumerate(t.rows):
        cells_info = []
        for c in r.cells:
            c_text = " // ".join(p.text.strip() for p in c.paragraphs if p.text.strip())
            cells_info.append(c_text)
        print(f"  Row {r_idx}: {cells_info}")

print('\n=== SECTIONS ===')
for s_idx, s in enumerate(doc.sections):
    print(f"Section {s_idx}: top={s.top_margin.pt}pt, bottom={s.bottom_margin.pt}pt, left={s.left_margin.pt}pt, right={s.right_margin.pt}pt, page_width={s.page_width.pt}pt, page_height={s.page_height.pt}pt")
    header_texts = [p.text for p in s.header.paragraphs if p.text.strip()]
    footer_texts = [p.text for p in s.footer.paragraphs if p.text.strip()]
    print(f'  Header: {header_texts}')
    print(f'  Footer: {footer_texts}')
