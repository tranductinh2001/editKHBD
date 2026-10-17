import sys
import io
import docx

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

doc = docx.Document('Mau KHBD THCS.docx')
print('=== FIRST 16 PARAGRAPHS OF Mau KHBD THCS.docx ===')
for i in range(min(16, len(doc.paragraphs))):
    p = doc.paragraphs[i]
    print(f"P{i:02d} [align={p.alignment}]: {repr(p.text)}")
    for r_idx, r in enumerate(p.runs):
        print(f"   run {r_idx}: text={repr(r.text)}, sz={r.font.size.pt if r.font.size else None}, b={r.font.bold}, i={r.font.italic}, u={r.font.underline}")

print('\n=== TABLE 0 in Mau KHBD THCS.docx ===')
if len(doc.tables) > 0:
    t0 = doc.tables[0]
    print(f"Table 0: {len(t0.rows)}x{len(t0.columns)}")
    for r_idx, r in enumerate(t0.rows):
        for c_idx, c in enumerate(r.cells):
            print(f"  Cell ({r_idx}, {c_idx}):")
            for p in c.paragraphs:
                print(f"     P: {repr(p.text)}")
                for r in p.runs:
                    print(f"        run: {repr(r.text)}, sz={r.font.size.pt if r.font.size else None}, b={r.font.bold}, i={r.font.italic}")
