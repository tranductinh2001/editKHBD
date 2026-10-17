import sys
import io
import docx

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

doc = docx.Document('Mau KHBD THCS.docx')
t = doc.tables[0]
print('Table alignment:', t.alignment)
tblPr = t._tbl.tblPr
tblBorders = tblPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tblBorders')
print('tblBorders:', tblBorders.xml if tblBorders is not None else "No tblBorders")

for r_idx, r in enumerate(t.rows):
    for c_idx, c in enumerate(r.cells):
        print(f'Cell ({r_idx}, {c_idx}): width={c.width.pt if c.width else None}pt')
        for p in c.paragraphs:
            print(f'   p: {repr(p.text)}')
            for run in p.runs:
                print(f'      run: {repr(run.text)}, sz={run.font.size.pt if run.font.size else None}, bold={run.font.bold}, italic={run.font.italic}')

print("\n=== Section 0 Header & Footer Details ===")
s = doc.sections[0]
print("Header paragraphs:")
for p in s.header.paragraphs:
    print(f"  P [align={p.alignment}]: {repr(p.text)}")
    for r in p.runs:
        print(f"     run: {repr(r.text)}, sz={r.font.size.pt if r.font.size else None}, b={r.font.bold}, i={r.font.italic}")

print("Footer paragraphs:")
for p in s.footer.paragraphs:
    print(f"  P [align={p.alignment}]: {repr(p.text)}")
    for r in p.runs:
        print(f"     run: {repr(r.text)}, sz={r.font.size.pt if r.font.size else None}, b={r.font.bold}, i={r.font.italic}")
