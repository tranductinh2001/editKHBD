import sys
import io
import os
import re
import docx
from docx.shared import Pt, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

def transform_to_thcs_template(file_path, output_path, config):
    doc = docx.Document(file_path)
    
    # 1. Page Margins & Header/Footer (Exact match to Mau KHBD THCS.docx)
    for s in doc.sections:
        s.top_margin = Pt(config.get("margin_top_pt", 56.7))
        s.bottom_margin = Pt(config.get("margin_bottom_pt", 56.7))
        s.left_margin = Pt(config.get("margin_left_pt", 56.7))
        s.right_margin = Pt(config.get("margin_right_pt", 42.55))
        
        # Header: Kế hoạch bài dạy {subject} {classes}, năm học {school_year}
        if config.get("enable_header", True):
            header = s.header
            for p in header.paragraphs:
                p.text = ""
            hp = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
            hp.text = f"Kế hoạch bài dạy {config.get('subject', 'Tin học')} {config.get('classes', '9')}, năm học {config.get('school_year', '2025-2026')}"
            hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if hp.runs:
                hp.runs[0].font.name = "Times New Roman"
                hp.runs[0].font.size = Pt(11)
                hp.runs[0].font.italic = True
                
        # Footer: {teacher_name} - {school_name}
        if config.get("enable_footer", True):
            footer = s.footer
            for p in footer.paragraphs:
                p.text = ""
            fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            fp.text = f"{config.get('teacher_name', 'Lê Thị Phương Thảo')} - {config.get('school_name', 'Trường TH&THCS Đăk Trôi')}"
            fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if fp.runs:
                fp.runs[0].font.name = "Times New Roman"
                fp.runs[0].font.size = Pt(10)
                fp.runs[0].font.italic = True

    # 2. Identify elements to remove before "I. Mục tiêu"
    body = doc.element.body
    elements_to_remove = []
    found_target = False
    for child in list(body):
        if isinstance(child, CT_P):
            p = Paragraph(child, doc)
            text_strip = p.text.strip().lower()
            if any(text_strip.startswith(x) for x in ["i. mục tiêu", "i. muc tieu", "1. mục tiêu", "i.mục tiêu"]) or (text_strip.startswith("i.") and "mục tiêu" in text_strip):
                found_target = True
                break
            else:
                elements_to_remove.append(child)
        elif isinstance(child, CT_Tbl):
            if not found_target:
                elements_to_remove.append(child)
                
    for el in elements_to_remove:
        body.remove(el)

    # 3. Create standardized elements (Mau KHBD THCS structure)
    new_elements = []
    
    # 3.1 TABLE 0 (1x2): AT THE VERY TOP
    table = doc.add_table(rows=1, cols=2)
    table.alignment = None
    table.autofit = False
    
    table.rows[0].cells[0].width = Pt(274.75)
    table.rows[0].cells[1].width = Pt(205.55)
    
    c0 = table.cell(0, 0)
    c1 = table.cell(0, 1)
    
    week = config.get("week", "1")
    periods = config.get("periods", "1, 2")
    teaching_date = config.get("teaching_date", "../../2026")
    teacher_name = config.get("teacher_name", "Lê Thị Phương Thảo")
    department = config.get("department", "Tổ Tự Nhiên")
    
    # Cell 0
    p_tuan = c0.paragraphs[0]
    p_tuan.paragraph_format.line_spacing = 1.15
    p_tuan.paragraph_format.space_after = Pt(2)
    run = p_tuan.add_run(f"Tuần:\t{week}")
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.italic = True
    
    p_tiet = c0.add_paragraph()
    p_tiet.paragraph_format.line_spacing = 1.15
    p_tiet.paragraph_format.space_after = Pt(2)
    run = p_tiet.add_run(f"Tiết: {periods}")
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.italic = True
    
    p_date = c0.add_paragraph()
    p_date.paragraph_format.line_spacing = 1.15
    p_date.paragraph_format.space_after = Pt(2)
    run = p_date.add_run(f"Ngày dạy: {teaching_date}")
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.italic = True

    # Cell 1
    p_gv = c1.paragraphs[0]
    p_gv.paragraph_format.line_spacing = 1.15
    p_gv.paragraph_format.space_after = Pt(2)
    run = p_gv.add_run(f"Họ và tên giáo viên: {teacher_name}")
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.italic = True
    
    p_to = c1.add_paragraph()
    p_to.paragraph_format.line_spacing = 1.15
    p_to.paragraph_format.space_after = Pt(2)
    run = p_to.add_run(f"Tổ chuyên môn: {department}")
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.italic = True
    
    new_elements.append(table._tbl)
    
    # Blank separator after table
    p_sep = doc.add_paragraph()
    p_sep.paragraph_format.space_after = Pt(4)
    new_elements.append(p_sep._p)

    # 3.2 TÊN BÀI DẠY: ...
    bai_num = config.get("bai_num", "1")
    bai_name = config.get("bai_name", "THẾ GIỚI KĨ THUẬT SỐ")
    if config.get("title_prefix_ten_bai_day", True):
        if bai_num and not f"bài {bai_num}".lower() in bai_name.lower():
            bai_title_full = f"TÊN BÀI DẠY: BÀI {bai_num}. {bai_name.upper()}"
        else:
            bai_title_full = f"TÊN BÀI DẠY: {bai_name.upper()}"
    else:
        bai_title_full = f"BÀI {bai_num}: {bai_name.upper()}" if bai_num else bai_name.upper()
        
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(2)
    p_title.paragraph_format.line_spacing = 1.15
    run = p_title.add_run(bai_title_full)
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.bold = True
    new_elements.append(p_title._p)

    # 3.3 Môn học/Hoạt động giáo dục: ...; lớp: ...
    subject = config.get("subject", "Tin học")
    classes = config.get("classes", "9")
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(2)
    p_sub.paragraph_format.line_spacing = 1.15
    run = p_sub.add_run(f"Môn học/Hoạt động giáo dục: {subject}; lớp: {classes}")
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    new_elements.append(p_sub._p)

    # 3.4 Thời gian thực hiện: (số tiết)
    duration = config.get("duration", "2")
    p_dur = doc.add_paragraph()
    p_dur.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_dur.paragraph_format.space_before = Pt(0)
    p_dur.paragraph_format.space_after = Pt(6)
    p_dur.paragraph_format.line_spacing = 1.15
    run = p_dur.add_run(f"Thời gian thực hiện: {duration} tiết")
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    new_elements.append(p_dur._p)

    for el in reversed(new_elements):
        body.remove(el)
        body.insert(0, el)
        
    # 4. Approval block at end
    if config.get("enable_approval_block", True):
        last_texts = [p.text.lower() for p in doc.paragraphs[-5:] if p.text.strip()]
        has_approval = any("duyệt của chuyên môn" in t or "duyệt của tổ chuyên môn" in t for t in last_texts)
        if not has_approval:
            p_app = doc.add_paragraph()
            p_app.paragraph_format.space_before = Pt(20)
            p_app.paragraph_format.space_after = Pt(10)
            p_app.paragraph_format.line_spacing = 1.15
            run1 = p_app.add_run("Duyệt của Chuyên môn\t\t\t\t        Duyệt của tổ chuyên môn")
            run1.font.name = "Times New Roman"
            run1.font.size = Pt(14)
            run1.font.bold = True

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Saved: {output_path}")

test_cfg = {
    "teacher_name": "Lê Thị Phương Thảo",
    "school_name": "Trường TH&THCS Đăk Trôi",
    "department": "Tổ Tự Nhiên",
    "school_year": "2025-2026",
    "subject": "Tin học",
    "classes": "9",
    "week": "1",
    "periods": "1, 2",
    "teaching_date": "16/9/2026",
    "bai_num": "1",
    "bai_name": "Thế giới kĩ thuật số",
    "duration": "2",
    "title_prefix_ten_bai_day": True,
    "enable_header": True,
    "enable_footer": True,
    "enable_approval_block": True
}

transform_to_thcs_template("Tin 9 bài 1.docx", "test_output_thcs/test_thcs_1.docx", test_cfg)
